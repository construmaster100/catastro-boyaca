"""
Scraper de los estudios economicos de la Camara de Comercio de Tunja (https://cctunja.org.co/estudios-economicos/).

Recorre las 12 secciones de estudios economicos, descarga cada documento publicado y genera:

    docs/camara_comercio/<seccion>/<archivo>        copia local de cada PDF / Excel
    docs/camara_comercio/catalogo_camara_comercio.xlsx
        hojas: Resumen (por seccion), Documentos, Indices (tabla de contenido de cada PDF), Tableros (Power BI)
    visor/camara/catalogo.json                       el mismo catalogo, para la interfaz web (camara.html)
    visor/camara/basemun.json, baseboy.json, hojas/  datos de la herramienta Excel de Boyaca en Cifras 2024
    visor/camara/analisis.html                       informe de analisis_cct.py

La extraccion de texto y tablas a Excel la hace extraer_cct.py (proceso mas largo).
Solo descarga lo que falta: se puede ejecutar de nuevo cuando la Camara publique documentos nuevos.

Uso: python scraper_cct.py
"""
import html as html_lib
import json
import re
import unicodedata
import warnings
from datetime import date
from pathlib import Path
from urllib.parse import unquote

import openpyxl
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ_WEB = "https://cctunja.org.co/estudios-economicos/"
CARPETA = Path(__file__).parent
DOCS = CARPETA / "docs" / "camara_comercio"
WEB = CARPETA / "visor" / "camara"
UA = {"User-Agent": "Mozilla/5.0 (catastro-boyaca; investigacion academica)"}
SECCIONES = {
    "boyaca-en-cifras": "Boyacá en Cifras",
    "boletin-economico": "Boletín económico",
    "dinamica-economica": "Dinámica económica",
    "tejido-empresarial": "Tejido empresarial",
    "censo-empresarial": "Censo empresarial",
    "concepto-economico": "Concepto económico",
    "notas-economicas": "Notas económicas",
    "estudios-de-tendencia": "Estudios de tendencia",
    "estudios-de-impacto-y-percepcion": "Estudios de impacto y percepción",
    "encuesta-ritmo-empresarial": "Encuesta de ritmo empresarial",
    "estimacion-del-potencial-comerciantes": "Estimación del potencial de comerciantes",
    "publicaciones-especiales": "Publicaciones especiales",
}
warnings.filterwarnings("ignore")


def limpio(t):
    return re.sub(r"\s+", " ", html_lib.unescape(re.sub(r"<[^>]+>", " ", t or ""))).strip()


def nombre_archivo(url):
    n = unquote(url.rsplit("/", 1)[-1])
    n = unicodedata.normalize("NFKD", n).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9._-]+", "_", n)


def titulo_de(url, texto):
    texto = limpio(texto)
    if texto and len(texto) > 6 and not re.match(r"(?i)^(descargar|ver|aqu[ií]|clic|pdf|leer m[aá]s)", texto):
        return texto[:160]
    n = re.sub(r"[-_]+", " ", unquote(url.rsplit("/", 1)[-1]).rsplit(".", 1)[0]).strip()
    return n[:1].upper() + n[1:]


def anio_de(url, titulo):
    anios = re.findall(r"(?<!\d)(20[0-3]\d)(?!\d)", titulo + " " + unquote(url.rsplit("/", 1)[-1]))
    if anios:
        return max(anios)
    m = re.search(r"/uploads/(20\d\d)/", url)
    return m.group(1) if m else ""


def indice_pdf(ruta):
    """Paginas y tabla de contenido (marcadores del PDF o lineas '....... N' de las primeras paginas)."""
    import pypdf
    r = pypdf.PdfReader(str(ruta))
    paginas, indice = len(r.pages), []

    def recorrer(items, nivel=0):
        for it in items:
            if isinstance(it, list):
                recorrer(it, nivel + 1)
            else:
                try:
                    indice.append({"t": str(it.title).strip(), "p": r.get_destination_page_number(it) + 1, "n": nivel})
                except Exception:
                    pass
    try:
        recorrer(r.outline)
    except Exception:
        pass
    if not indice:
        patron = re.compile(r"^\s*((?:\d+(?:\.\d+)*\.?\s+)?[A-ZÁÉÍÓÚÑ][^\n]{3,110}?)\s*\.{4,}\s*(\d{1,3})\s*$")
        for i in range(min(12, paginas)):
            for linea in (r.pages[i].extract_text() or "").splitlines():
                m = patron.match(linea)
                if m and 0 < int(m.group(2)) <= paginas:
                    t = m.group(1).strip()
                    indice.append({"t": t, "p": int(m.group(2)), "n": min(t.split(" ")[0].count(".") if t[:1].isdigit() else 0, 2)})
    vistos, unicos = set(), []
    for e in indice:
        if (e["t"], e["p"]) not in vistos:
            vistos.add((e["t"], e["p"]))
            unicos.append(e)
    return paginas, unicos[:250]


def excel_boyaca_en_cifras(ruta):
    """Convierte la herramienta Excel de Boyaca en Cifras a los JSON que usa la interfaz."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    (WEB / "hojas").mkdir(parents=True, exist_ok=True)
    for ws in wb.worksheets:
        filas = [[(round(v, 4) if isinstance(v, float) else v) for v in f] for f in ws.iter_rows(values_only=True)]
        filas = [f for f in filas if any(v is not None for v in f)]
        (WEB / "hojas" / f"{nombre_archivo(ws.title)}.json").write_text(
            json.dumps({"hoja": ws.title, "filas": filas}, ensure_ascii=False, default=str), encoding="utf-8")
    f = [list(x) for x in wb["BaseMun"].iter_rows(values_only=True)]
    grupos, grupo = [], ""
    for g in f[0]:
        grupo = str(g).strip() if g else grupo
        grupos.append(grupo)
    columnas = [{"k": str(c).strip(), "grupo": grupos[i]} for i, c in enumerate(f[1])]
    filas = [[(round(v, 3) if isinstance(v, float) else v) for v in r] for r in f[2:] if r[0]]
    (WEB / "basemun.json").write_text(json.dumps({"columnas": columnas, "filas": filas}, ensure_ascii=False), encoding="utf-8")
    f = [list(x) for x in wb["BaseBoy"].iter_rows(values_only=True)]
    cols = [str(c).strip() for c in f[0]]
    deps = {r[0]: {cols[i]: (round(v, 4) if isinstance(v, float) else v) for i, v in enumerate(r) if i} for r in f[1:] if r[0]}
    (WEB / "baseboy.json").write_text(json.dumps({"columnas": cols[1:], "departamentos": deps}, ensure_ascii=False), encoding="utf-8")


def hoja_excel(wb, nombre, cabecera, filas, anchos=None):
    ws = wb.create_sheet(nombre)
    ws.append(cabecera)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F4D2E")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for f in filas:
        ws.append(f)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, a in enumerate(anchos or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    return ws


def main():
    DOCS.mkdir(parents=True, exist_ok=True)
    WEB.mkdir(parents=True, exist_ok=True)
    catalogo, tableros, descripciones = [], [], {}
    for slug, nombre in SECCIONES.items():
        url_sec = RAIZ_WEB + slug + "/"
        print(f"== {nombre}", flush=True)
        pagina = requests.get(url_sec, timeout=60, headers=UA).text
        cuerpo = pagina[pagina.find("<main"):] if "<main" in pagina else pagina
        descripciones[slug] = [p for p in (limpio(x) for x in re.findall(r"<p[^>]*>(.*?)</p>", cuerpo, re.S)) if len(p) > 60][:4]
        for u in sorted(set(re.findall(r"https://app\.powerbi\.com/view\?r=[^\"'&<\s]+", cuerpo))):
            tableros.append({"seccion": nombre, "slug": slug, "titulo": f"Tablero Power BI · {nombre}", "url": html_lib.unescape(u)})
        vistos = set()
        carpeta = DOCS / slug
        carpeta.mkdir(exist_ok=True)
        for url, texto in re.findall(r'<a[^>]+href="([^"]+\.(?:pdf|xlsx?|csv|docx?))"[^>]*>(.*?)</a>', cuerpo, re.I | re.S):
            url = html_lib.unescape(url).replace("http://", "https://")
            if url in vistos or "Auditoria" in url:
                continue
            vistos.add(url)
            archivo = nombre_archivo(url)
            ruta = carpeta / archivo
            if not ruta.exists():
                try:
                    r = requests.get(url, timeout=300, headers=UA)
                    r.raise_for_status()
                    ruta.write_bytes(r.content)
                    print(f"   descargado {archivo} ({len(r.content) / 1e6:.1f} MB)", flush=True)
                except Exception as e:
                    print(f"   ERROR {archivo}: {e}", flush=True)
                    continue
            titulo = titulo_de(url, texto)
            tipo = "Excel" if archivo.lower().endswith((".xlsx", ".xls", ".csv")) else "Word" if ".doc" in archivo.lower() else "PDF"
            doc = {"seccion": nombre, "slug": slug, "titulo": titulo, "tipo": tipo, "anio": anio_de(url, titulo),
                   "archivo": f"../docs/camara_comercio/{slug}/{archivo}", "url": url, "mb": round(ruta.stat().st_size / 1e6, 2)}
            if tipo == "PDF":
                try:
                    doc["paginas"], doc["indice"] = indice_pdf(ruta)
                except Exception as e:
                    doc["paginas"], doc["indice"] = None, []
                    print(f"   aviso: no se pudo leer {archivo}: {e}")
            elif slug == "boyaca-en-cifras" and tipo == "Excel":
                excel_boyaca_en_cifras(ruta)
            catalogo.append(doc)
        n = sum(1 for d in catalogo if d["slug"] == slug)
        print(f"   {n} documentos", flush=True)

    orden = list(SECCIONES)
    catalogo.sort(key=lambda d: (orden.index(d["slug"]), -int(d["anio"] or 0), d["titulo"]))
    (WEB / "catalogo.json").write_text(json.dumps({
        "fuente": "Cámara de Comercio de Tunja", "pagina": RAIZ_WEB, "consultado": date.today().isoformat(),
        "secciones": [{"slug": s, "nombre": SECCIONES[s], "url": RAIZ_WEB + s + "/", "descripcion": descripciones.get(s, [])}
                      for s in orden],
        "documentos": catalogo, "tableros": tableros}, ensure_ascii=False, indent=1), encoding="utf-8")

    # ---- catalogo en Excel
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    resumen = []
    for s in orden:
        ds = [d for d in catalogo if d["slug"] == s]
        anios = sorted({d["anio"] for d in ds if d["anio"]})
        resumen.append([SECCIONES[s], len(ds), sum(d.get("paginas") or 0 for d in ds), round(sum(d["mb"] for d in ds), 1),
                        f"{anios[0]}–{anios[-1]}" if anios else "", RAIZ_WEB + s + "/"])
    resumen.append(["TOTAL", len(catalogo), sum(d.get("paginas") or 0 for d in catalogo), round(sum(d["mb"] for d in catalogo), 1), "", ""])
    ws = hoja_excel(wb, "Resumen", ["Sección", "Documentos", "Páginas", "MB", "Años", "Página web"], resumen, [42, 12, 10, 10, 12, 70])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    hoja_excel(wb, "Documentos", ["Sección", "Año", "Título", "Tipo", "Páginas", "MB", "Entradas de índice", "Archivo local", "URL original"],
               [[d["seccion"], d["anio"], d["titulo"], d["tipo"], d.get("paginas"), d["mb"], len(d.get("indice") or []),
                 "docs/camara_comercio/" + d["archivo"].split("docs/camara_comercio/")[-1], d["url"]] for d in catalogo],
               [30, 8, 70, 8, 9, 8, 10, 60, 80])
    hoja_excel(wb, "Indices", ["Sección", "Documento", "Año", "Nivel", "Entrada", "Página"],
               [[d["seccion"], d["titulo"], d["anio"], e.get("n", 0), e["t"], e["p"]] for d in catalogo for e in d.get("indice") or []],
               [30, 50, 8, 7, 80, 8])
    hoja_excel(wb, "Tableros", ["Sección", "Título", "URL"], [[t["seccion"], t["titulo"], t["url"]] for t in tableros], [30, 45, 120])
    wb.save(DOCS / "catalogo_camara_comercio.xlsx")

    md = CARPETA / "informes" / "analisis_boyaca_en_cifras.md"
    if md.exists():
        try:
            import markdown
            (WEB / "analisis.html").write_text(markdown.markdown(md.read_text(encoding="utf-8"), extensions=["tables"]), encoding="utf-8")
        except ImportError:
            pass
    print(f"Listo: {len(catalogo)} documentos ({sum(d['mb'] for d in catalogo):,.0f} MB) en docs/camara_comercio/ "
          f"y catalogo_camara_comercio.xlsx; {len(tableros)} tableros")


if __name__ == "__main__":
    main()
