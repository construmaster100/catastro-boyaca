"""
Exporta los datos propios del proyecto a Excel en docs/datos/:

    docs/datos/municipios_boyaca.xlsx
        Municipios     123 municipios: areas, % urbano/rural, jerarquia urbana, indicadores CCT 2024, uso del suelo
        Provincias     13 provincias con sus totales
        Jerarquia      jerarquia urbana ordenada
        Coincidencias  cruce pagina CCT <-> repositorio (codigo DANE y nombre) con predios y economia
        Fuentes        origen de cada grupo de columnas

Uso: python exportar_excel.py   (despues de exportar_geo.py y analisis_cct.py)
"""
import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

CARPETA = Path(__file__).parent
SALIDA = CARPETA / "docs" / "datos"


def formatear(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F4D2E")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 60)), 1):
        largo = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 45)


def main():
    SALIDA.mkdir(parents=True, exist_ok=True)
    mun = pd.read_csv(CARPETA / "geo" / "boyaca_municipios_puntos.csv", dtype={"codigo": str})
    prov = pd.DataFrame([f["properties"] for f in json.loads(
        (CARPETA / "geo" / "boyaca_provincias.geojson").read_text(encoding="utf-8"))["features"]])
    jer = mun[["jerarquia_urbana", "codigo", "municipio", "provincia", "categoria_urbana", "area_urbana_ha", "pct_urbano",
               "pct_rural", "area_km2", "pct_de_boyaca", "poblacion", "densidad_hab_km2", "valor_agregado"]] \
        .sort_values("jerarquia_urbana")
    coin = pd.read_csv(CARPETA / "informes" / "coincidencias_municipios.csv", sep=";", decimal=",", dtype={"codigo": str})
    fuentes = pd.DataFrame([
        ["Predios, áreas de predios, uso del suelo", "IGAC · Dato_Fundamental_Catastro (U_TERRENO, R_TERRENO)", "catastro_colombia.py"],
        ["Límites municipales", "IGAC · ordenamientoterritorial/pendientepromediomunicipio", "exportar_web.py"],
        ["Zonas urbanas, áreas, jerarquía urbana", "Cálculo propio sobre predios IGAC", "estadisticas_urbanas.py"],
        ["Provincia, población, economía, social, agua", "Cámara de Comercio de Tunja · Boyacá en Cifras 2024", "importar_cct.py"],
        ["Uso del suelo: urbano = actividad económica; rural 10–300 m² = vivienda; > 300 m² = agropecuario",
         "Criterio del usuario aplicado a los predios IGAC", "exportar_geo.py"],
    ], columns=["Datos", "Fuente", "Script"])
    ruta = SALIDA / "municipios_boyaca.xlsx"
    with pd.ExcelWriter(ruta, engine="openpyxl") as xl:
        for nombre, df in [("Municipios", mun), ("Provincias", prov), ("Jerarquia", jer),
                           ("Coincidencias", coin), ("Fuentes", fuentes)]:
            df.to_excel(xl, sheet_name=nombre, index=False)
            formatear(xl.sheets[nombre])
    print(f"Listo: {ruta.relative_to(CARPETA)} ({len(mun)} municipios, {len(prov)} provincias)")


if __name__ == "__main__":
    main()
